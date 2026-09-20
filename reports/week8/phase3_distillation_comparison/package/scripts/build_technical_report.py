#!/usr/bin/env python3
"""Assemble the eight-chapter evidence report from reviewed repository sources."""
from pathlib import Path
import re
import shutil
from common import ROOT, sha256, write_json

REPORTS = ROOT/'reports'


def selected(path, numbers, chapter):
    text = (ROOT/path).read_text()
    parts = re.split(r'(?m)^## ', text)[1:]
    result = []
    index = 0
    for part in parts:
        match = re.match(r'(\d+)\.', part)
        if not match or int(match[1]) not in numbers:
            continue
        index += 1
        heading, _, content = part.partition('\n')
        heading = re.sub(r'^\d+\.\s*', '', heading)
        # Existing nested numbering belonged to the weekly report, not this report.
        content = re.sub(r'(?m)^### (?:\d+[.\d]*\s*)?', '#### ', content)
        content = re.sub(r'(?m)^#### (?:\d+[.\d]*\s*)?', '#### ', content)
        # Remove embedded old diagrams; this report inserts verified figures below.
        content = re.sub(r'!\[[^\]]*\]\([^)]*\)', '', content)
        # Source-local links stay human-readable; chapter source is in references.
        content = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', content)
        result.append(f'### {chapter}.{index} {heading}\n\n{content.strip()}')
    return '\n\n'.join(result)


def assemble():
    sources = [
        'deliverables/week1/day3/REPORT.md',
        'deliverables/week2/day10/REPORT.md',
        'deliverables/week3/day16/REPORT.md',
        'Submission/Week4/Day21_Weekly_Report_and_Final_Model_Archive/Week4_DPO_Preference_Alignment_Report.md',
        'deliverables/week5/day27/REPORT.md',
        'deliverables/week6/day33/REPORT.md',
        'Submission/Week7/量化对比报告.md',
        'Submission/Week7/第7周部署与应用报告.md',
    ]
    text = '# Qwen大模型实习综合技术报告\n\n'
    text += '2026年9月18日阶段证据整合稿\n\n'
    text += '本文面向实习导师和后续复现者，整合前七周的环境、数据、训练、评估和部署证据，并说明第八周自动化代码与尚待执行的蒸馏实验。前七周已有完整阶段报告，主要归档缺口在主入口与跨周总结；模型质量方面仍保留第二周、第五周未达标和第六周端到端可靠性不足的真实结论。第八周已完成冻结数据准备、正式SFT与DPO训练和合并、三模型完整CEval/CMMLU评测及Gemini评分；确定性诊断发现数学和代码短板。两份匿名人工评分已完成并汇总；蒸馏、部署监督器实测仍未完成，本文是阶段报告，不是全部验收通过的终稿。\n\n'
    text += '阅读说明：第1至6章以已冻结周报为依据重整，保留实验当时的数据与限制；章节中提及的Day和source文件属于原始归档，完整路径见文末对应资料。第7至8章描述本轮实现、验证和后续工作。历史周报所写“提交”在本综合稿中仅表示已有可提交文件，未核验教师实际接收。\n\n'
    text += '## 第1章 环境与架构\n\n'
    text += '历史实验使用NVIDIA RTX3090的24GB显存、Python3.10.20、PyTorch2.5.1与CUDA12.1，主要文本模型为Qwen2.5-7B-Instruct。选择7B模型是为了在单卡预算内保留足够表达能力，训练采用QLoRA以降低权重与优化器开销。macOS本机承担脚本开发、证据整理和报告生成，不能把本地Python可执行等同于CUDA训练环境已验证。\n\n'
    text += selected(sources[0], [1,2,4,5,6,9], 1)
    text += '\n\n## 第2章 数据工程\n\n'
    text += '数据工程先定义来源、抽样规模和格式，再执行清洗与去重。历史结果保持冻结，第八周新增训练验证划分属于新的数据产物，二者不混用。\n\n'
    text += selected(sources[1], [3,4,7,9], 2)
    text += '\n\n## 第3章 SFT优化实验\n\n'
    text += selected(sources[2], [3,4,5,6,7,9,10,11,12,13,14], 3)
    text += '\n\n![图1 第1周两轮身份微调的真实TensorBoard训练曲线](figures/identity_loss.png)\n\n'
    text += '图1是第1周历史训练截图，用于展示训练监控方式；不是第3周九组实验的曲线。九组实验数值以各自日志和第三周报告为准。\n\n'
    text += '![图2 第3周固定20题的模型五维评分雷达图](figures/model_radar.png)\n\n'
    text += '图2来自第三周冻结的人类评分底表；不得与第八周AI裁判结果混称同一种评估。\n\n'
    text += '## 第4章 DPO偏好对齐\n\n'
    text += selected(sources[3], [2,3,4,5,6,7], 4)
    text += '\n\n![图3 第4周最终纠偏DPO的Rewards曲线](figures/dpo_rewards.png)\n\n'
    text += '图3对应最终四十步纠偏模型。初始二百四十步实验不符合Rejected趋势字面要求，保留为诊断证据。\n\n'
    text += '## 第5章 多模态与Agent实践\n\n'
    text += selected(sources[4], [3,4,5,6,7,8,10,11], 5)
    text += '\n\n![图4 第20层目标文字对图像区域的注意力分布](figures/vlm_attention.png)\n\n'
    text += '图4为第5周场景案例的真实热力图。Qwen2-VL此处实际提取self-attention中的视觉token区域，不能虚构独立cross-attention模块；该实现偏差已有方法说明。\n\n'
    text += '### 5.9 Agent工具与端到端表现\n\n'
    # Strip duplicated 5.x section labels for the combined VLM/Agent chapter.
    agent = selected(sources[5], [2,3,4,5,6,7,8], 5)
    agent = re.sub(r'(?m)^### 5\.\d+ ', '#### ', agent)
    text += agent
    text += '\n\n## 第6章 部署与量化\n\n'
    for source in sources[6:]:
        contents = (ROOT/source).read_text().split('\n',1)[1]
        contents = re.sub(r'(?m)^## ', '### ', contents)
        contents = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', contents)
        # A compact table for portrait pages; weight sizes remain in original report.
        if '量化对比' in source:
            contents = re.sub(r'(?ms)^\| 模型 \|.*?(?=\n\n)', '| 模型 | 峰值显存MiB | 显存下降 | tokens/s | PPL |\n|---|---:|---:|---:|---:|\n| FP16 | 17622 | 基准 | 41.2438 | 7.55026 |\n| AWQ | 8286 | 52.9792% | 22.4332 | 8.18261 |\n| GPTQ | 8290 | 52.9565% | 20.8910 | 8.02914 |', contents, count=1)
        text += contents+'\n\n'
    text += (REPORTS/'week8/report_supplement.md').read_text()+'\n\n'
    text += '## 参考资料与证据索引\n\n'
    for i, source in enumerate(sources,1):
        title=(ROOT/source).read_text().splitlines()[0].lstrip('# ')
        text += f'[{i}] [{title}](../{source})。本仓库历史实验与报告，2026年9月14日核对。文件：{source}\n\n'
    text += '[9] [OpenCompass官方快速开始](https://opencompass.readthedocs.io/en/stable/get_started/quick_start.html)。用于核对CLI评估入口。\n\n'
    text += '[10] [Transformers知识蒸馏教程](https://huggingface.co/docs/transformers/v4.48.1/en/tasks/knowledge_distillation_for_image_classification)。用于解释软标签、温度与KL项，与本项目序列级方法区分。\n\n'
    text += '[11] 本仓库最新版《实习需求》，源文件指纹见week8/requirements_source.json；本周任务拆解见docs/week8_execution_plan.md。\n'
    (REPORTS/'technical_report.md').write_text(text)
    prefix = text.split('## 第5章')[0]
    counts = {'chapters_1_to_4_chinese_characters': len(re.findall('[\u4e00-\u9fff]',prefix)), 'total_chinese_characters': len(re.findall('[\u4e00-\u9fff]',text)), 'chapter_count': len(re.findall(r'(?m)^## 第[1-8]章',text)), 'status': 'stage_report_training_eval_human_complete_distillation_deployment_pending'}
    if counts['chapters_1_to_4_chinese_characters'] < 3000 or counts['total_chinese_characters'] < 6000:
        raise ValueError(counts)
    write_json(REPORTS/'week8/report_validation.json', counts)
    write_json(REPORTS/'week8/report_sources.json', {p: sha256(ROOT/p) for p in sources})
    figures = {
        'identity_loss.png': 'deliverables/week1/day5/evidence/tensorboard_loss.png',
        'model_radar.png': 'deliverables/week3/day14/source/results/model_scores_radar.png',
        'dpo_rewards.png': 'Submission/Week4/Day21_Weekly_Report_and_Final_Model_Archive/DPO_Rewards_Curves.png',
        'vlm_attention.png': 'deliverables/week5/day24/source/results/heatmaps/case-01-scene-mountain-triptych.png',
    }
    for target, source in figures.items():
        shutil.copy2(ROOT/source, REPORTS/'figures'/target)
    return text


def build_docx(text):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from PIL import Image, ImageDraw, ImageFont
    canvas = Image.new('RGB', (1500, 370), 'white'); draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 24)
    labels = ['Archive data', 'Clean / split', 'SFT + DPO', 'Eval + judge', 'API + Gradio']
    for i, label in enumerate(labels):
        x=20+i*295
        draw.rounded_rectangle((x,70,x+260,165), radius=12, fill='#e7f0f5', outline='#38627c', width=3)
        draw.text((x+18,105), label, fill='#17384f', font=font)
        if i<4:
            draw.line((x+265,117,x+291,117),fill='#38627c',width=4)
    draw.text((40,225), 'Each stage: config + inputs + hashes + logs + explicit completion status',fill='#17384f',font=font)
    draw.text((40,275), 'Quick mode: local data smoke + archived score replay. No fresh inference.',fill='#17384f',font=font)
    canvas.save(REPORTS/'figures/pipeline.png')
    doc = Document(); section = doc.sections[0]
    section.top_margin=Inches(.75);section.bottom_margin=Inches(.75)
    section.left_margin=Inches(.75);section.right_margin=Inches(.75)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Arial Unicode MS';style._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
        style.font.color.rgb=RGBColor.from_string('000000')
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(24)
    for element in list(doc.styles['Title']._element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    doc.styles['Heading 1'].font.size=Pt(17)
    doc.styles['Heading 2'].font.size=Pt(13)
    footer=section.footer.paragraphs[0];footer.alignment=2
    footer.add_run('Qwen 实习综合报告  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    def plain(value):
        value=re.sub(r'\[([^\]]+)\]\([^)]*\)',r'\1',value)
        return value.replace('`','').replace('**','')
    lines=text.splitlines();i=0;code=False;table_number=0;last_heading='实验记录'
    while i<len(lines):
        line=lines[i];i+=1
        if not line.strip():continue
        if line.startswith('```'):
            if code and doc.paragraphs:
                doc.paragraphs[-1].paragraph_format.keep_with_next=False
            code=not code;continue
        if code:
            p=doc.add_paragraph(line);p.paragraph_format.line_spacing=1
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.keep_with_next=True
            for r in p.runs:r.font.size=Pt(8)
            continue
        image=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        if image:
            path=REPORTS/image[2]
            with Image.open(path) as img:w,h=img.size
            width=min(6.55, 5.5*w/h)
            doc.add_picture(str(path),width=Inches(width))
            p=doc.add_paragraph(image[1]);p.alignment=1
            for r in p.runs:r.font.size=Pt(9)
            continue
        if line.startswith('|'):
            table_number+=1
            caption=doc.add_paragraph(f'表{table_number} {last_heading}')
            caption.paragraph_format.keep_with_next=True
            block=[line]
            while i<len(lines) and lines[i].startswith('|'):
                block.append(lines[i]);i+=1
            rows=[[plain(v.strip()) for v in x.strip('|').split('|')] for x in block if not re.match(r'^\|[\s:|\-]+$',x)]
            table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
            for j,value in enumerate(rows[0]):table.rows[0].cells[j].text=value
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows[1:]:
                cells=table.add_row().cells
                for j,value in enumerate(row[:len(cells)]):cells[j].text=value
            for row in table.rows:
                cant_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant_split)
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(2)
                        if len(table.rows) <= 18:
                            para.paragraph_format.keep_with_next=True
                        for run in para.runs:
                            run.font.size=Pt(8)
                            run.font.name='Arial Unicode MS'
                            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
            for cell in table.rows[-1].cells:
                for para in cell.paragraphs:para.paragraph_format.keep_with_next=False
            continue
        if line.startswith('# '):doc.add_paragraph(plain(line[2:]),style='Title')
        elif line.startswith('## '):
            heading=doc.add_heading(plain(line[3:]),level=1)
            heading.paragraph_format.page_break_before=True
        elif line.startswith('### '):
            last_heading=plain(line[4:])
            doc.add_heading(last_heading,level=2)
        elif line.startswith('#### '):doc.add_heading(plain(line[5:]),level=3)
        else:doc.add_paragraph(plain(line))
    doc.save(REPORTS/'technical_report.docx')


if __name__ == '__main__':
    build_docx(assemble())
